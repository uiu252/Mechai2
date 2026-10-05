"""Read both Student Data sections; preserve source values and provenance."""
from pathlib import Path
from io import BytesIO
import hashlib
import json
import math
import re
from openpyxl import load_workbook
from schemas import Product, CATEGORIES

DEFAULT_WORKBOOK = Path(__file__).parent / 'data' / 'Halloween_Line_Review_Case_Study_Data.xlsx'

class WorkbookError(ValueError):
    pass

def key(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())

ALIASES = {
    'item': 'item', 'category': 'category', 'requiredfacings': 'required_facings',
    'retailprice': 'retail_price', 'proposedretail': 'retail_price', 'unitcost': 'unit_cost',
    'grossmargin': 'gross_margin_dollars', 'grossmargindollars': 'gross_margin_dollars',
    'potentialgmbasedonpurchasequantity': 'gross_margin_dollars',
    'grossmarginrate': 'gross_margin_rate', 'estimatedunitmarginrate': 'gross_margin_rate',
    'fullpricesellthrough': 'full_price_sell_through', 'starrating': 'star_rating',
    'reviews': 'reviews', 'mandatorypurchasequantity': 'mandatory_purchase_quantity',
    'mandatorypurchasequanity': 'mandatory_purchase_quantity', 'licensed': 'licensed',
    'merchantnotes': 'merchant_notes', 'vendorsays': 'vendor_says', 'currentchannel': 'placement',
    'unitsbought': 'units_bought', 'unitssoldfullprice': 'units_sold_full_price',
    'unitssoldclearance': 'units_sold_clearance', 'storeunits': 'store_units',
    'onlineunits': 'online_units', 'totalsales': 'total_sales', 'gmperfacing': 'historical_gm_per_facing',
}
TEXT = {'item', 'category', 'merchant_notes', 'vendor_says', 'placement'}
INTS = {'required_facings', 'reviews', 'mandatory_purchase_quantity', 'units_bought',
        'units_sold_full_price', 'units_sold_clearance', 'store_units', 'online_units'}

def normalize_category(value):
    norm = key(str(value).replace('&', 'and'))
    for category in CATEGORIES:
        if key(category.replace('&', 'and')) == norm:
            return category
    return str(value).strip()

def normalize_placement(value):
    placements = {'instore': 'In Store', 'onlineonly': 'Online Only', 'removed': 'Removed'}
    if key(value) not in placements:
        raise WorkbookError(f'Unknown placement: {value!r}')
    return placements[key(value)]

def value_for(field, value):
    if value is None or str(value).strip().lower() in ('', 'n/a', 'na', 'not available', '-'):
        return None
    if isinstance(value, str) and value.startswith('='):
        raise WorkbookError('Formula lacks a cached value. Recalculate and save the workbook in Excel.')
    if field == 'licensed':
        if key(value) in ('yes', 'true', '1'): return True
        if key(value) in ('no', 'false', '0'): return False
        raise WorkbookError(f'Invalid Licensed? value: {value!r}')
    if field in TEXT: return str(value).strip()
    raw = str(value).strip().replace(',', '').replace('$', '')
    negative = raw.startswith('(') and raw.endswith(')')
    raw = raw.strip('()')
    percent = raw.endswith('%')
    number = float(raw.rstrip('%')) * (-1 if negative else 1) / (100 if percent else 1)
    if not math.isfinite(number): raise WorkbookError('Non-finite numeric value.')
    if field in INTS:
        if not number.is_integer(): raise WorkbookError(f'{field} must be an integer.')
        return int(number)
    return number

def load_products(source=DEFAULT_WORKBOOK):
    try:
        raw = source if isinstance(source, bytes) else (source.read() if hasattr(source, 'read') else Path(source).read_bytes())
        if len(raw) > 15_000_000: raise WorkbookError('Workbook must be less than 15 MB.')
        workbook = load_workbook(BytesIO(raw), data_only=True, read_only=True)
        if 'Student Data' not in workbook.sheetnames:
            raise WorkbookError('The workbook must include a Student Data worksheet.')
        result, mapping, source_type = [], None, None
        for row_number, row in enumerate(workbook['Student Data'].iter_rows(values_only=True), 1):
            headers = [key(c) for c in row]
            if 'item' in headers and 'category' in headers and 'requiredfacings' in headers:
                mapping = {i: ALIASES[h] for i, h in enumerate(headers) if h in ALIASES}
                source_type = 'new_candidate' if any('mandatorypurchase' in h for h in headers) else 'current'
                continue
            if not mapping: continue
            item_index = next(i for i, f in mapping.items() if f == 'item')
            category_index = next(i for i, f in mapping.items() if f == 'category')
            if row[item_index] is None or row[category_index] is None: continue
            try:
                record = {field: value_for(field, row[i]) for i, field in mapping.items()}
                record['category'] = normalize_category(record['category'])
                record['source_type'] = source_type
                record['source_row'] = row_number
                record['placement'] = normalize_placement(record.get('placement') or 'Removed')
                record['original_placement'] = record['placement']
                if source_type == 'current' and re.search(r'\blicensed\b', record['item'], re.I):
                    record['licensed'] = True
                result.append(Product.model_validate(record))
            except Exception as exc:
                raise WorkbookError(f'Student Data row {row_number}: {exc}') from exc
        workbook.close()
        if not result: raise WorkbookError('No product rows found in Student Data.')
        names = [p.item.casefold() for p in result]
        if len(names) != len(set(names)): raise WorkbookError('Duplicate item names are not allowed.')
        if {p.source_type for p in result} != {'current', 'new_candidate'}:
            raise WorkbookError('Both current assortment and new-candidate sections are required.')
        return result
    except WorkbookError: raise
    except Exception as exc: raise WorkbookError(f'Cannot read workbook: {exc}') from exc

def fingerprint(products):
    serialized = [p.model_dump() for p in sorted(products, key=lambda p: p.item)]
    return hashlib.sha256(json.dumps(serialized, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

def source_warnings(products):
    warnings = []
    for p in products:
        if p.category not in CATEGORIES:
            warnings.append(f'{p.item}: unrecognized category {p.category!r}.')
        missing = [f.replace('_', ' ') for f in ('retail_price', 'unit_cost', 'gross_margin_dollars', 'gross_margin_rate') if getattr(p, f) is None]
        if missing: warnings.append(f'{p.item}: Not available — {", ".join(missing)}.')
    reaper = next((p for p in products if p.item == '7 ft Animated Reaper'), None)
    sitcom = next((p for p in products if p.item == '7 ft Licensed Sitcom Character'), None)
    if reaper and sitcom and reaper.required_facings != sitcom.required_facings:
        warnings.append(f'Supplier note calls the Sitcom Character the same footprint as the Reaper, but Required Facings are {sitcom.required_facings} and {reaper.required_facings}. Validation uses the numeric Required Facings column.')
    return warnings
